"""Workbook scanning: real exports do not put a header in cell A1.

Each case here is a shape that a POS or accounting export actually produces and
that the first implementation rejected outright.
"""

from __future__ import annotations

import io

import pandas as pd
import pytest
from openpyxl import Workbook

from app.semantic.engine import map_dataframe
from app.services.parsing import FileError, parse_upload
from app.services.profiling import profile_dataframe
from app.services.workbook import extract_table

COLUMNS = ["Bill No", "Date", "Branch", "Item", "Qty", "Amount", "Discount"]


def sales_rows(n=40):
    return [
        [f"B{i}", f"2026-01-{(i % 28) + 1:02d}", "Jubilee Hills", "Biryani", 2, 320, 10]
        for i in range(1, n + 1)
    ]


def workbook_bytes(build) -> bytes:
    wb = Workbook()
    build(wb)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def frame_bytes(df: pd.DataFrame, sheets: dict | None = None) -> bytes:
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer) as writer:
        if sheets:
            for name, sheet in sheets.items():
                sheet.to_excel(writer, sheet_name=name, index=False)
        else:
            df.to_excel(writer, sheet_name="Sheet1", index=False)
    return buffer.getvalue()


def test_header_below_title_rows_is_found():
    """The most common real shape: company name, report title, blank, header."""
    def build(wb):
        ws = wb.active
        ws.title = "Sales"
        ws.append(["Spice Route Restaurants"])
        ws.append(["Monthly Sales Report - January 2026"])
        ws.append([])
        ws.append(COLUMNS)
        for row in sales_rows():
            ws.append(row)

    df, report = parse_upload("report.xlsx", workbook_bytes(build))
    assert list(df.columns) == COLUMNS
    assert len(df) == 40
    assert report.header_row == 4
    assert report.skipped_top_rows == 3
    assert any("row 4" in note for note in report.notes)


def test_merged_banner_does_not_become_the_header():
    def build(wb):
        ws = wb.active
        ws.merge_cells("A1:G1")
        ws["A1"] = "SALES REGISTER"
        ws.append([])
        ws.append([])
        ws.append(COLUMNS)
        for row in sales_rows():
            ws.append(row)

    df, _ = parse_upload("register.xlsx", workbook_bytes(build))
    assert list(df.columns) == COLUMNS


def test_data_sheet_is_chosen_over_a_cover_sheet():
    """Reading only the first sheet is why a valid workbook looked empty."""
    cover = pd.DataFrame({"Report": ["Sales export", "Generated 01-Feb-2026"]})
    data = pd.DataFrame(sales_rows(), columns=COLUMNS)
    content = frame_bytes(None, {"Cover": cover, "Data": data})

    df, report = parse_upload("export.xlsx", content)
    assert report.sheet_name == "Data"
    assert report.sheets_available == ["Cover", "Data"]
    assert len(df) == 40
    assert any("sheets" in note for note in report.notes)


def test_a_named_sheet_can_be_requested():
    cover = pd.DataFrame({"Report": ["Sales export"]})
    data = pd.DataFrame(sales_rows(), columns=COLUMNS)
    content = frame_bytes(None, {"Cover": cover, "Data": data})

    df, report = parse_upload("export.xlsx", content, sheet="Cover")
    assert report.sheet_name == "Cover"
    assert list(df.columns) == ["Report"]

    with pytest.raises(FileError, match="no sheet named"):
        parse_upload("export.xlsx", content, sheet="Nope")


def test_totals_row_is_not_counted_as_a_transaction():
    rows = sales_rows()
    frame = pd.DataFrame(rows, columns=COLUMNS)
    with_total = pd.concat(
        [frame, pd.DataFrame([{"Bill No": "TOTAL", "Amount": 12800, "Qty": 80}])]
    )

    df, report = parse_upload("sales.xlsx", frame_bytes(with_total))
    assert len(df) == 40
    assert report.dropped_total_rows == 1
    assert df["Amount"].sum() == 40 * 320


def test_repeated_headings_are_kept_distinct():
    frame = pd.DataFrame(sales_rows(), columns=COLUMNS)
    frame.columns = [*COLUMNS[:-1], "Amount"]  # two columns literally named Amount

    df, _ = parse_upload("dupes.xlsx", frame_bytes(frame))
    assert len(df.columns) == len(set(df.columns))
    assert sum(c.startswith("Amount") for c in df.columns) == 2


def test_spacer_column_is_dropped():
    frame = pd.DataFrame(sales_rows(), columns=COLUMNS)
    frame.insert(0, "", [None] * len(frame))

    df, report = parse_upload("spaced.xlsx", frame_bytes(frame))
    assert list(df.columns) == COLUMNS
    assert report.dropped_blank_columns >= 1


def test_xlsx_content_named_xls_is_still_read():
    """POS exports routinely save xlsx bytes under a .xls name."""
    frame = pd.DataFrame(sales_rows(), columns=COLUMNS)
    df, _ = parse_upload("export.xls", frame_bytes(frame))
    assert len(df) == 40


def test_spreadsheet_with_no_extension_is_recognised_by_content():
    frame = pd.DataFrame(sales_rows(), columns=COLUMNS)
    df, _ = parse_upload("export", frame_bytes(frame))
    assert len(df) == 40


def test_types_are_recovered_after_scanning():
    """The scanner reads raw cells; real dtypes must come back afterwards."""
    frame = pd.DataFrame(sales_rows(), columns=COLUMNS)
    df, _ = parse_upload("sales.xlsx", frame_bytes(frame))
    assert pd.api.types.is_numeric_dtype(df["Amount"])
    assert pd.api.types.is_datetime64_any_dtype(df["Date"])


def test_currency_text_is_read_as_a_number():
    frame = pd.DataFrame(sales_rows(), columns=COLUMNS)
    frame["Amount"] = frame["Amount"].map(lambda v: f"₹ {v:,.2f}")
    df, _ = parse_upload("sales.xlsx", frame_bytes(frame))
    assert pd.api.types.is_numeric_dtype(df["Amount"])
    assert df["Amount"].sum() == 40 * 320


@pytest.mark.parametrize(
    "build",
    [
        lambda wb: [wb.active.append(["Spice Route"]), wb.active.append([]),
                    wb.active.append(COLUMNS)] + [wb.active.append(r) for r in sales_rows()],
    ],
)
def test_awkward_files_still_classify_as_sales(build):
    """The point of scanning: detection has to work end to end, not just parse."""
    df, _ = parse_upload("messy.xlsx", workbook_bytes(build))
    _, signals = profile_dataframe(df)
    kind, proposals = map_dataframe(df, signals)
    assert kind == "sales"
    assert {p.concept for p in proposals} >= {"REVENUE", "BRANCH", "ORDER_ID"}


def test_csv_still_takes_the_simple_path():
    frame = pd.DataFrame(sales_rows(), columns=COLUMNS)
    df, report = extract_table(frame.to_csv(index=False).encode(), "sales.csv")
    assert report.sheet_name is None
    assert list(df.columns) == COLUMNS


def test_csv_with_title_rows_is_also_handled():
    body = "Spice Route\nJanuary 2026\n\n" + pd.DataFrame(
        sales_rows(), columns=COLUMNS
    ).to_csv(index=False)
    df, report = parse_upload("sales.csv", body.encode())
    assert list(df.columns) == COLUMNS
    assert report.header_row == 4
